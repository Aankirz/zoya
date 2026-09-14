"""Office files for document_agent (STACK §9 tier 1, D32): write, read back, edit, open.

Every format is read into one shape and written from it, so "create", "read row 3" and "make slide
3 shorter" share code: a document is a title plus parts (slides, or sections with lines), a table
is columns, rows and the columns to total. Editing rewrites the whole file from that shape, so only
files Zoya made (under ~/Documents/Zoya/) may be edited or shared.

Libraries (pinned in pyproject.toml): python-pptx https://python-pptx.readthedocs.io/en/latest/ ,
python-docx https://python-docx.readthedocs.io/en/latest/ , openpyxl (real `=SUM()` formulas; no
cached values until an app opens the file, so read-back totals are summed here)
https://openpyxl.readthedocs.io/en/stable/usage.html , reportlab (PDF with one outline entry per
section) https://docs.reportlab.com/reportlab/userguide/ch2_graphics/ ; PDFs are read back with
PDFKit https://developer.apple.com/documentation/pdfkit/pdfdocument . Opening uses the app's bundle
id (`open -b`, `man open`).

Anything read from a file is untrusted (§12.2): read-backs are wrapped for the model.
These tools must only run inside document_agent, so there is no TOOLS list here.
"""

from __future__ import annotations

import csv
import re
import subprocess
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from strands import tool

from zoya import safety
from zoya.config import (
    DOCUMENT_READ_ROOTS,
    DOCUMENTS_DIR,
    OFFICE_MAX_PARTS,
    OFFICE_MAX_ROWS,
    OFFICE_TEXT_MAX_CHARS,
    OPEN_APP_TIMEOUT_S,
)
from zoya.tools import ToolError

DOCUMENT_KINDS = {"pptx", "docx", "pdf", "md"}
TABLE_KINDS = {"xlsx", "csv"}
BUNDLES = {  # first installed wins
    "pptx": ("com.apple.Keynote", "com.microsoft.Powerpoint"),
    "docx": ("com.apple.Pages", "com.microsoft.Word"),
    "xlsx": ("com.apple.Numbers", "com.microsoft.Excel"),
    "csv": ("com.apple.Numbers", "com.microsoft.Excel"),
    "pdf": ("com.apple.Preview",),
    "md": ("com.apple.TextEdit",),
}
TOTAL_LABEL = "Total"
SLUG = re.compile(r"[^a-z0-9]+")
SLUG_MAX_CHARS = 48
PDF_MARGIN_PT = 60
PDF_LINE_PT = 16
PDF_HEADING_PT = 24
LIST_FILES_MAX = 15


@dataclass
class Doc:
    title: str
    parts: list[dict[str, Any]] = field(default_factory=list)  # {"title": str, "lines": [str]}


@dataclass
class Table:
    title: str
    columns: list[str]
    rows: list[list[Any]]
    total_columns: list[str] = field(default_factory=list)


# --- Paths ----------------------------------------------------------------------------------------


def _new_path(title: str, kind: str) -> Path:
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    stem = SLUG.sub("-", title.lower()).strip("-")[:SLUG_MAX_CHARS] or "document"
    path, number = DOCUMENTS_DIR / f"{stem}.{kind}", 2
    while path.exists():
        path, number = DOCUMENTS_DIR / f"{stem}-{number}.{kind}", number + 1
    return path


def zoya_file(path: str) -> Path:
    """A file Zoya made (inside ~/Documents/Zoya/ after resolving links): edit, open, share."""
    resolved = Path(path).expanduser().resolve()
    if not resolved.is_relative_to(DOCUMENTS_DIR.resolve()) or not resolved.is_file():
        raise ToolError(
            "I can only change or share files I made, in Documents, Zoya. (Call list_files.)"
        )
    return resolved


def readable_file(path: str) -> Path:
    resolved = Path(path).expanduser().resolve()
    if not any(resolved.is_relative_to(root.resolve()) for root in DOCUMENT_READ_ROOTS):
        raise ToolError("I can read files in your Documents, Downloads or Desktop folders.")
    if not resolved.is_file():
        raise ToolError("I can't find that file. (Call list_files and use a path from it.)")
    return resolved


def _kind(path: Path) -> str:
    kind = path.suffix.lower().lstrip(".")
    if kind not in DOCUMENT_KINDS | TABLE_KINDS:
        raise ToolError("I can work with slides, Word, PDF, Markdown, Excel and CSV files.")
    return kind


# --- Writers --------------------------------------------------------------------------------------


def write_document(doc: Doc, path: Path) -> None:
    {"pptx": _write_pptx, "docx": _write_docx, "pdf": _write_pdf, "md": _write_md}[_kind(path)](
        doc, path
    )


def _write_pptx(doc: Doc, path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    title_layout, content_layout = deck.slide_layouts[0], deck.slide_layouts[1]
    for number, part in enumerate(doc.parts):
        layout = title_layout if number == 0 and not part["lines"] else content_layout
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = part["title"]
        if layout is content_layout:
            body = slide.placeholders[1].text_frame
            body.text = part["lines"][0] if part["lines"] else ""
            for line in part["lines"][1:]:
                body.add_paragraph().text = line
    deck.core_properties.title = doc.title
    deck.save(path)


def _write_docx(doc: Doc, path: Path) -> None:
    from docx import Document

    word = Document()
    word.core_properties.title = doc.title
    word.add_heading(doc.title, level=0)
    for part in doc.parts:
        word.add_heading(part["title"], level=1)
        for line in part["lines"]:
            word.add_paragraph(line)
    word.save(path)


def _write_md(doc: Doc, path: Path) -> None:
    lines = [f"# {doc.title}", ""]
    for part in doc.parts:
        lines += [f"## {part['title']}", "", *part["lines"], ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_pdf(doc: Doc, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path), pagesize=A4)
    canvas.setTitle(doc.title)
    width, height = A4
    y = height - PDF_MARGIN_PT

    def line(text: str, size: int, key: str = "") -> None:
        nonlocal y
        if y < PDF_MARGIN_PT:
            canvas.showPage()
            y = height - PDF_MARGIN_PT
        if key:  # one outline entry per section: read back as the section list
            canvas.bookmarkPage(key, fit="XYZ", top=y + size)
            canvas.addOutlineEntry(text, key, level=0)
        canvas.setFont("Helvetica-Bold" if key or size > PDF_LINE_PT else "Helvetica", size)
        canvas.drawString(PDF_MARGIN_PT, y, text[: int((width - 2 * PDF_MARGIN_PT) / (size / 2))])
        y -= size * 1.5

    line(doc.title, PDF_HEADING_PT)
    for number, part in enumerate(doc.parts, start=1):
        line(part["title"], PDF_LINE_PT, key=f"section{number}")
        for text in part["lines"]:
            line(text, PDF_LINE_PT - 4)
    canvas.save()


def write_table(table: Table, path: Path) -> None:
    (_write_xlsx if _kind(path) == "xlsx" else _write_csv)(table, path)


def _write_xlsx(table: Table, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = table.title[:31] or "Sheet1"  # Excel's sheet-name limit
    sheet.append(table.columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in table.rows:
        sheet.append(row)
    if table.total_columns:
        last = len(table.rows) + 1
        totals = [TOTAL_LABEL] + [""] * (len(table.columns) - 1)
        for name in table.total_columns:
            letter = get_column_letter(table.columns.index(name) + 1)
            totals[table.columns.index(name)] = f"=SUM({letter}2:{letter}{last})"
        sheet.append(totals)
        for cell in sheet[last + 1]:
            cell.font = Font(bold=True)
    book.save(path)


def _write_csv(table: Table, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(table.columns)
        writer.writerows(table.rows)
        if table.total_columns:
            sums = column_totals(table)
            writer.writerow(
                [TOTAL_LABEL]
                + [_number_text(sums[c]) if c in sums else "" for c in table.columns[1:]]
            )


# --- Readers --------------------------------------------------------------------------------------


def read_document_file(path: Path) -> Doc:
    return {"pptx": _read_pptx, "docx": _read_docx, "pdf": _read_pdf, "md": _read_md}[_kind(path)](
        path
    )


def _read_pptx(path: Path) -> Doc:
    from pptx import Presentation

    deck = Presentation(str(path))
    parts = []
    for slide in list(deck.slides)[:OFFICE_MAX_PARTS]:
        title_shape = slide.shapes.title
        lines = [
            paragraph.text
            for shape in slide.shapes
            if shape.has_text_frame and shape != title_shape
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        ]
        parts.append({"title": title_shape.text if title_shape is not None else "", "lines": lines})
    return Doc(deck.core_properties.title or path.stem, parts)


def _read_docx(path: Path) -> Doc:
    from docx import Document

    word = Document(str(path))
    doc = Doc(word.core_properties.title or path.stem)
    for paragraph in word.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style == "Title":
            doc.title = paragraph.text
        elif style.startswith("Heading"):
            doc.parts.append({"title": paragraph.text, "lines": []})
        elif paragraph.text.strip():
            if not doc.parts:
                doc.parts.append({"title": "", "lines": []})
            doc.parts[-1]["lines"].append(paragraph.text)
    return doc


def _read_md(path: Path) -> Doc:
    doc = Doc(path.stem)
    for raw in path.read_text(encoding="utf-8").splitlines():
        text = raw.strip()
        if text.startswith("# "):
            doc.title = text[2:]
        elif text.startswith("#"):
            doc.parts.append({"title": text.lstrip("#").strip(), "lines": []})
        elif text:
            if not doc.parts:
                doc.parts.append({"title": "", "lines": []})
            doc.parts[-1]["lines"].append(text)
    return doc


def _read_pdf(path: Path) -> Doc:
    """Sections from the PDF outline, their text between one heading and the next."""
    from Foundation import NSURL
    from Quartz import PDFDocument

    pdf = PDFDocument.alloc().initWithURL_(NSURL.fileURLWithPath_(str(path)))
    if pdf is None or pdf.isLocked():
        raise ToolError("I can't open that PDF. It may be password protected.")
    text = str(pdf.string() or "")
    attributes = pdf.documentAttributes() or {}
    outline = pdf.outlineRoot()
    headings = [
        str(outline.childAtIndex_(i).label())
        for i in range(outline.numberOfChildren() if outline else 0)
    ]
    doc = Doc(str(attributes.get("Title") or path.stem))
    rest = text
    for number, heading in enumerate(headings):
        start = rest.find(heading)
        if start < 0:
            continue
        rest = rest[start + len(heading) :]
        following = headings[number + 1] if number + 1 < len(headings) else None
        body = rest[: rest.find(following)] if following and following in rest else rest
        doc.parts.append(
            {"title": heading, "lines": [ln for ln in body.splitlines() if ln.strip()]}
        )
    if not headings:
        doc.parts.append({"title": "", "lines": [ln for ln in text.splitlines() if ln.strip()]})
    return doc


def read_table_file(path: Path) -> Table:
    if _kind(path) == "xlsx":
        return _read_xlsx(path)
    with path.open(newline="", encoding="utf-8") as handle:
        rows = [row for _, row in zip(range(OFFICE_MAX_ROWS + 2), csv.reader(handle), strict=False)]
    return _table_from_rows(path.stem, rows, totals_from_row=True)


def _read_xlsx(path: Path) -> Table:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True).active
    rows = [
        ["" if value is None else value for value in row]
        for _, row in zip(
            range(OFFICE_MAX_ROWS + 2), sheet.iter_rows(values_only=True), strict=False
        )
    ]
    return _table_from_rows(sheet.title, rows, totals_from_row=True)


def _table_from_rows(title: str, rows: list[list[Any]], totals_from_row: bool) -> Table:
    if not rows:
        return Table(title, [], [])
    columns = [str(value) for value in rows[0]]
    body = rows[1:]
    total_columns: list[str] = []
    if totals_from_row and body and str(body[-1][0]).strip() == TOTAL_LABEL:
        total_row = body.pop()
        total_columns = [
            columns[i] for i, value in enumerate(total_row) if i and str(value).strip()
        ]
    return Table(title, columns, [list(row) for row in body], total_columns)


# --- Numbers --------------------------------------------------------------------------------------


def to_number(value: Any) -> Decimal | None:
    """Parser (tested): 1200, "1,200", "₹1,200.50", "Rs 300" → Decimal; anything else → None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    cleaned = re.sub(r"(?i)^\s*(?:₹|rs\.?|inr|\$)\s*", "", str(value)).replace(",", "").strip()
    try:
        return Decimal(cleaned) if cleaned else None
    except InvalidOperation:
        return None


def column_totals(table: Table) -> dict[str, Decimal]:
    totals = {}
    for name in table.total_columns:
        index = table.columns.index(name)
        cells = [to_number(row[index]) for row in table.rows if index < len(row)]
        totals[name] = sum((cell for cell in cells if cell is not None), Decimal(0))
    return totals


def _number_text(value: Decimal) -> str:
    return f"{value:,.2f}".rstrip("0").rstrip(".")


# --- Spoken read-backs (parsers: tested) ----------------------------------------------------------


def document_summary(doc: Doc, noun: str) -> str:
    titles = [part["title"] or f"{noun} {n}" for n, part in enumerate(doc.parts, start=1)]
    count = f"{len(doc.parts)} {noun}{'' if len(doc.parts) == 1 else 's'}"
    return f"{doc.title}: {count}" + (f", {', '.join(titles)}." if titles else ".")


def table_summary(table: Table) -> str:
    rows = len(table.rows)
    parts = [
        f"{table.title}: columns {', '.join(table.columns)}; {rows} row{'' if rows == 1 else 's'}"
    ]
    parts += [f"total {name} {_number_text(value)}" for name, value in column_totals(table).items()]
    return "; ".join(parts) + "."


def part_text(doc: Doc, number: int, noun: str) -> str:
    if not 1 <= number <= len(doc.parts):
        raise ToolError(f"There are {len(doc.parts)} {noun}s. Which one?")
    part = doc.parts[number - 1]
    return f"{noun.capitalize()} {number}: {part['title']}. " + " ".join(part["lines"])


def row_text(table: Table, number: int) -> str:
    """Row 1 is the first data row, as a person counts it (the header isn't a row)."""
    if not 1 <= number <= len(table.rows):
        raise ToolError(f"There are {len(table.rows)} rows. Which one?")
    row = table.rows[number - 1]
    cells = [
        f"{name} {'' if i >= len(row) else row[i]}".strip() for i, name in enumerate(table.columns)
    ]
    return f"Row {number}: " + ", ".join(cells) + "."


NOUNS = {"pptx": "slide", "docx": "section", "pdf": "section", "md": "section"}


def _speak(text: str) -> str:
    return safety.wrap_untrusted(text[:OFFICE_TEXT_MAX_CHARS])


def summarize(path: Path) -> str:
    kind = _kind(path)
    if kind in TABLE_KINDS:
        return _speak(table_summary(read_table_file(path)))
    return _speak(document_summary(read_document_file(path), NOUNS[kind]))


# --- Tools ----------------------------------------------------------------------------------------


def _clean_parts(parts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not parts or len(parts) > OFFICE_MAX_PARTS:
        raise ToolError(f"Give between 1 and {OFFICE_MAX_PARTS} parts.")
    return [
        {"title": str(p.get("title", "")).strip(), "lines": [str(x) for x in p.get("lines", [])]}
        for p in parts
    ]


@tool
def create_document(title: str, kind: str, parts: list[dict[str, Any]]) -> str:
    """Create slides, a Word document, a PDF or a Markdown file in ~/Documents/Zoya/ and return
    its path and structure.

    Args:
        title: Document title, e.g. "Renewable Energy".
        kind: "pptx" (slides), "docx" (Word), "pdf" or "md".
        parts: Slides or sections in order, each {"title": str, "lines": [str]}. For slides the
            lines are bullets; a first slide with no lines is a title slide.
    """
    if kind not in DOCUMENT_KINDS:
        raise ToolError("Use kind pptx, docx, pdf or md.")
    path = _new_path(title, kind)
    write_document(Doc(title.strip(), _clean_parts(parts)), path)
    return f"Saved {path}.\n{summarize(path)}"


@tool
def create_table(
    title: str, kind: str, columns: list[str], rows: list[list[Any]], total_columns: list[str]
) -> str:
    """Create an Excel (.xlsx, with real SUM formulas on a Total row) or CSV sheet in
    ~/Documents/Zoya/ and return its path, columns, row count and totals.

    Args:
        title: Sheet title, e.g. "Monthly Expenses".
        kind: "xlsx" or "csv".
        columns: Header names, e.g. ["Category", "Amount (rupees)"].
        rows: Data rows in column order; put amounts as plain numbers, e.g. ["Rent", 15000].
        total_columns: Columns to total, e.g. ["Amount (rupees)"]; [] for no total row.
    """
    if kind not in TABLE_KINDS:
        raise ToolError("Use kind xlsx or csv.")
    table = _checked_table(title, columns, rows, total_columns)
    path = _new_path(title, kind)
    write_table(table, path)
    return f"Saved {path}.\n{summarize(path)}"


def _checked_table(
    title: str, columns: list[str], rows: list[list[Any]], total_columns: list[str]
) -> Table:
    if not columns or len(rows) > OFFICE_MAX_ROWS:
        raise ToolError(f"Give column names and at most {OFFICE_MAX_ROWS} rows.")
    unknown = [name for name in total_columns if name not in columns]
    if unknown or columns[0] in total_columns:
        raise ToolError("Totals must name columns other than the first one.")
    width = len(columns)
    return Table(
        title.strip(), [str(c) for c in columns], [list(r)[:width] for r in rows], total_columns
    )


@tool
def read_file_structure(path: str) -> str:
    """Read back a document's structure to speak: title and slide/section names, or a sheet's
    columns, row count and totals.

    Args:
        path: Full path, e.g. one returned by create_document.
    """
    return summarize(readable_file(path))


@tool
def read_file_part(path: str, number: int) -> str:
    """Read one slide, section or row (row 1 = first data row) to speak.

    Args:
        path: Full path of the file.
        number: Which slide, section or row, counting from 1.
    """
    file = readable_file(path)
    kind = _kind(file)
    if kind in TABLE_KINDS:
        return _speak(row_text(read_table_file(file), number))
    return _speak(part_text(read_document_file(file), number, NOUNS[kind]))


@tool
def edit_file_part(path: str, number: int, title: str, lines: list[str]) -> str:
    """Replace one slide/section (title + lines) or one row (lines = cell values) of a file Zoya
    made; number = count + 1 adds a new one at the end; lines = [] with title "DELETE" removes it.

    Args:
        path: Full path of a file in ~/Documents/Zoya/.
        number: Which slide, section or row, counting from 1.
        title: New slide/section title ("" for rows).
        lines: New bullets/paragraphs, or the row's cell values in column order.
    """
    file = zoya_file(path)
    if _kind(file) in TABLE_KINDS:
        table = read_table_file(file)
        table.rows = replace_item(table.rows, number, title, list(lines)[: len(table.columns)])
        write_table(table, file)
    else:
        doc = read_document_file(file)
        doc.parts = replace_item(doc.parts, number, title, {"title": title, "lines": list(lines)})
        write_document(doc, file)
    return f"Updated {file}.\n{summarize(file)}"


def replace_item(items: list[Any], number: int, title: str, new: Any) -> list[Any]:
    """Parser (tested): replace item N, append at N = len + 1, delete when title is "DELETE"."""
    if not 1 <= number <= len(items) + 1:
        raise ToolError(f"There are {len(items)}. Pick one of those, or {len(items) + 1} to add.")
    if title == "DELETE":
        if number > len(items):
            raise ToolError("There's nothing there to remove.")
        return items[: number - 1] + items[number:]
    return items[: number - 1] + [new] + items[number:]


@tool
def list_files() -> str:
    """List the files Zoya made in ~/Documents/Zoya/, newest first, with their full paths."""
    if not DOCUMENTS_DIR.is_dir():
        return "No files yet."
    files = sorted(
        (f for f in DOCUMENTS_DIR.iterdir() if f.suffix.lower().lstrip(".") in BUNDLES),
        key=lambda f: f.stat().st_mtime,
        reverse=True,
    )[:LIST_FILES_MAX]
    return safety.wrap_untrusted("\n".join(str(f) for f in files)) if files else "No files yet."


@tool
def open_file(path: str) -> str:
    """Open a file Zoya made in its Mac app (Keynote, Pages, Numbers, Preview, TextEdit). Only
    when the user asks to open it.

    Args:
        path: Full path of a file in ~/Documents/Zoya/.
    """
    file = zoya_file(path)
    for bundle in BUNDLES[_kind(file)]:
        if _installed(bundle):
            try:
                subprocess.run(
                    ["open", "-b", bundle, str(file)], check=True, timeout=OPEN_APP_TIMEOUT_S
                )
            except (subprocess.SubprocessError, OSError) as error:
                raise ToolError("I couldn't open that file.") from error
            return f"Opened {file.name}."
    raise ToolError("No app on this Mac opens that kind of file.")


def _installed(bundle: str) -> bool:
    from AppKit import NSWorkspace

    return NSWorkspace.sharedWorkspace().URLForApplicationWithBundleIdentifier_(bundle) is not None


FILE_TOOLS = [
    create_document,
    create_table,
    read_file_structure,
    read_file_part,
    edit_file_part,
    list_files,
    open_file,
]
